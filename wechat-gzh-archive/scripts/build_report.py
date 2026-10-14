#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
组装「公众号文章汇总表」: 序号/发布时间/标题/内容分类(多标签)/类型/访问链接/是否抓取到/
阅读量/点赞/在看数/评论数/分享数/存档位置 -> xlsx + csv。
用法: python build_report.py <articles.json> <存档目录> [metrics.json] [-o 输出文件名]
- articles.json 与 scrape_content.py 用同一套「标题归一化去重」，保证序号与 .md 文件名前缀对应。
- metrics.json(可选): {content_url: {read_num,old_like_num,like_num,comment_count,share_num}} (次幂 data2 结果)。
- 内容分类为关键词启发式(标题为主+开头500字为辅), 边界样例可能需人工复核。
依赖: openpyxl (xlsx); csv 用标准库。
"""
import os, re, sys, json, csv, glob
from collections import Counter

# ---------- 分类词表 ----------
AD_EXPLICIT = ["赞助", "商单", "恰饭", "邀约", "特约", "金主", "品牌合作", "付费合作", "brand推广", "广告合作"]
AD_PROMO = ["送30天", "送7天", "送会员", "免费领", "限时免费", "0元", "薅羊毛", "优惠券", "官网直降", "限量名额", "内测名额", "扫码领取"]
AD_BRAND = "豆包工作"   # 示例: 集中推广的商业产品即商单系列, 使用时按目标号实际替换/增补
REVIEW_T = ["实测", "测评", "评测", "横评", "上手", "开箱"]
REVIEW_B = ["实测", "测评", "评测", "横评"]
CATS = {
    "干货教程": ["教程", "指南", "手册", "方法", "如何", "怎么", "工作流", "技巧", "步骤", "从0到1", "入门", "搭建", "打造", "提示词", "SOP", "实战", "全流程", "最佳实践", "攻略", "用法", "精通", "掌握", "姿势", "跑通"],
    "产品发布": ["做了个", "做了", "造了", "搞了个", "上线了", "开源", "全新", "正式版", "2.0", "V2", "新版", "发布了", "上线"],
    "资讯观点": ["解读", "报告", "分析", "观点", "思考", "我认为", "怎么看", "聊聊", "谈谈", "趋势", "预测", "泄露", "官宣", "突破", "拿下", "刷屏", "炸了", "意味着", "启示", "评价", "竞品"],
    "周刊合集": ["周刊", "合集", "盘点", "清单", "推荐", "导航", "汇总", "精选", "榜单", "资源", "合辑", "个最"],
    "个人随笔": ["随笔", "感悟", "年度", "年终", "复盘", "冒泡", "碎碎念", "感想", "心路", "焦虑", "一点思考"],
}
ORDER = ["接广告", "产品测评", "干货教程", "产品发布", "资讯观点", "周刊合集", "个人随笔"]


def _cnt(text, kws):
    return sum(len(re.findall(re.escape(k), text, re.I)) for k in kws)


def classify(title, head):
    tags = []
    if (_cnt(title, AD_EXPLICIT) or AD_BRAND in title or _cnt(title, AD_PROMO)
            or _cnt(head, AD_EXPLICIT) or _cnt(head, AD_PROMO) >= 2):
        tags.append("接广告")
    if _cnt(title, REVIEW_T) or _cnt(head, REVIEW_B) >= 2:
        tags.append("产品测评")
    for cat in ["干货教程", "产品发布", "资讯观点", "周刊合集", "个人随笔"]:
        if _cnt(title, CATS[cat]) >= 1 or _cnt(head, CATS[cat]) >= 2:
            tags.append(cat)
    if not tags:
        tags = ["其他"]
    return "，".join([t for t in ORDER if t in tags] + [t for t in tags if t not in ORDER])


def norm_title(t):
    t = re.sub(r"<[^>]+>", "", t or "")
    t = re.sub(r"\s+", "", t)
    return t[:40]


def detect_type(fp):
    try:
        return "图文" if "本文为小绿书/图片消息" in open(fp, encoding="utf-8").read(500) else "文章"
    except Exception:
        return "文章"


def clean(t):
    return (t or "").replace("\n", " ").replace("|", "/").strip()


def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    src, mddir = sys.argv[1], sys.argv[2]
    metrics_f = None
    out_name = "文章汇总表"
    args = sys.argv[3:]
    i = 0
    while i < len(args):
        if args[i] == "-o":
            out_name = args[i + 1]; i += 2
        else:
            metrics_f = args[i]; i += 1
    metrics = json.load(open(metrics_f, encoding="utf-8")) if metrics_f and os.path.exists(metrics_f) else {}

    raw = json.load(open(src, encoding="utf-8"))
    seen, arts = set(), []
    for a in raw:
        k = norm_title(a.get("title"))
        if k in seen:
            continue
        seen.add(k); arts.append(a)

    prefix2file = {}
    for fp in glob.glob(os.path.join(mddir, "*.md")):
        fn = os.path.basename(fp)
        if fn in ("index.md",) or fn.startswith("_"):
            continue
        m = re.match(r"(\d{3})_", fn)
        if m:
            prefix2file[int(m.group(1))] = fn

    def mval(url, key):
        d = metrics.get(url)
        if not d or "_error" in d:
            return "—"
        v = d.get(key)
        return v if v is not None else 0

    HEADER = ["序号", "发布时间", "标题", "内容分类", "类型", "访问链接", "是否抓取到",
              "阅读量", "点赞", "在看数", "评论数", "分享数", "存档位置(.md文件名)"]
    rows = []
    for i, a in enumerate(arts, 1):
        url = a["content_url"]
        fn = prefix2file.get(i)
        if fn:
            status, loc = "是", fn
            typ = detect_type(os.path.join(mddir, fn))
            head = open(os.path.join(mddir, fn), encoding="utf-8").read(900)
            head = re.sub(r"^---.*?---\s*", "", head, flags=re.S)[:500]
        else:
            status, loc, typ, head = "否", "—", "—", ""
        cat = classify((a.get("title") or ""), head)
        rows.append([i, (a.get("published_at") or "").replace("T", " "), clean(a.get("title")),
                     cat, typ, url, status,
                     mval(url, "read_num"), mval(url, "old_like_num"), mval(url, "like_num"),
                     mval(url, "comment_count"), mval(url, "share_num"), loc])

    # CSV
    with open(out_name + ".csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(HEADER); w.writerows(rows)
    print("CSV:", out_name + ".csv")

    # XLSX
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "文章汇总"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F6BED")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for idx, wd in enumerate([6, 19, 42, 22, 7, 46, 10, 9, 8, 9, 9, 9, 34], 1):
        ws.column_dimensions[chr(64 + idx)].width = wd
    green = PatternFill("solid", fgColor="E2EFDA"); red = PatternFill("solid", fgColor="FCE4E4")
    imgfill = PatternFill("solid", fgColor="FFF2CC"); adfill = PatternFill("solid", fgColor="F8CBAD")
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(vertical="center")
        row[3].alignment = Alignment(wrap_text=True, vertical="center")
        if "接广告" in str(row[3].value):
            row[3].fill = adfill; row[3].font = Font(bold=True, color="9C0006")
        row[4].alignment = Alignment(horizontal="center", vertical="center")
        if row[4].value == "图文":
            row[4].fill = imgfill
        cell = row[5]
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = cell.value; cell.font = Font(color="0563C1", underline="single")
        row[6].fill = green if row[6].value == "是" else red
        row[6].alignment = Alignment(horizontal="center", vertical="center")
        for c in row[7:12]:
            c.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(c.value, int):
                c.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{ws.max_row}"
    wb.save(out_name + ".xlsx")
    print("XLSX:", out_name + ".xlsx")
    cc = Counter()
    for r in rows:
        for t in str(r[3]).split("，"):
            cc[t] += 1
    have = sum(1 for r in rows if isinstance(r[7], int))
    print(f"\n共 {len(rows)} 篇 | 有指标 {have} | 分类分布: " + ", ".join(f"{k}{v}" for k, v in cc.most_common()))


if __name__ == "__main__":
    main()
